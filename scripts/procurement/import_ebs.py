"""
One-time import of EBS D.O.I. (BRUT).xlsx into procurement_erp table in elefant-erp.db.

When you get a new export from ERP, run this again — it replaces the table.

Usage: python -m procurement.import_ebs [--file path/to/EBS.xlsx]
"""
import argparse
import sqlite3
import time
import openpyxl
from procurement.config import DB_PATH


EBS_COLUMNS = [
    ('sku',            'TEXT'),
    ('ean',            'TEXT'),
    ('article_code',   'TEXT'),
    ('article_desc',   'TEXT'),
    ('author',         'TEXT'),
    ('supplier',       'TEXT'),
    ('rrp',            'REAL'),
    ('discount',       'REAL'),
    ('stock_online',   'REAL'),
    ('avail_rec',      'REAL'),
    ('sales_l1w',      'REAL'),
    ('sales_lm',       'REAL'),
    ('sales_l2m',      'REAL'),
    ('sales_ls',       'REAL'),
    ('sales_ly',       'REAL'),
    ('publish_date',   'TEXT'),
    ('category',       'TEXT'),
    ('subcategory',    'TEXT'),
    ('manufacturer',   'TEXT'),
    ('supplier_code',  'TEXT'),
    ('total_reserved', 'REAL'),
    ('avail_mm_auchan','REAL'),
    ('collection',     'TEXT'),
    ('format',         'TEXT'),
    ('age_group',      'TEXT'),
    ('sku_abon',       'TEXT'),
    ('purchases_all',  'REAL'),
    ('sales_all',      'REAL'),
    ('out_of_print',   'TEXT'),
    ('unavailable',    'TEXT'),
    ('avail_custf',    'REAL'),
    ('avail_stpr',     'REAL'),
    ('purchase_price', 'REAL'),
]


def import_ebs(ebs_path, db_path=DB_PATH):
    t0 = time.time()

    print(f"Reading {ebs_path}...")
    wb = openpyxl.load_workbook(ebs_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    print(f"  {len(rows)} rows ({time.time()-t0:.1f}s)")

    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA journal_mode=WAL")

    conn.execute("DROP TABLE IF EXISTS procurement_erp")
    cols_def = ', '.join(f"{name} {typ}" for name, typ in EBS_COLUMNS)
    conn.execute(f"CREATE TABLE procurement_erp ({cols_def})")

    placeholders = ', '.join('?' for _ in EBS_COLUMNS)
    n_cols = len(EBS_COLUMNS)

    def clean_row(row):
        r = list(row[:n_cols])
        # Normalize publish_date to ISO string
        val = r[15]  # publish_date
        if val is not None:
            from datetime import datetime
            if isinstance(val, datetime):
                r[15] = val.strftime('%Y-%m-%d')
            else:
                r[15] = str(val)[:10]
        return r

    conn.executemany(
        f"INSERT INTO procurement_erp VALUES ({placeholders})",
        (clean_row(r) for r in rows)
    )

    conn.execute("CREATE INDEX IF NOT EXISTS idx_erp_article ON procurement_erp(article_code)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_erp_sku ON procurement_erp(sku)")

    conn.commit()
    cnt = conn.execute("SELECT count(*) FROM procurement_erp").fetchone()[0]
    conn.close()

    print(f"  Imported {cnt:,} rows into procurement_erp ({time.time()-t0:.1f}s)")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--file', default='arhivă/EBS D.O.I. (BRUT).xlsx')
    parser.add_argument('--db', default=DB_PATH)
    args = parser.parse_args()
    import_ebs(args.file, args.db)


if __name__ == '__main__':
    main()
